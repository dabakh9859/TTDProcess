"""Data loading & preprocessing for the AI sidecar.

Mirrors the logic of `ai-poc/data_loader.py` but lives inside the sidecar so
the runtime is self-contained once packaged. The `load_toa5_xlsx` helper is
the same parser used during the POC; we keep the pickle-gzip cache for
fast reloads.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

EXCEL_EPOCH = pd.Timestamp("1899-12-30")

# Z-score clip for training/inference windows. Some processed channels (e.g.
# T600) are near-constant (std ~0.1) with rare dropout artifacts, which become
# enormous z-scores (|z| ~ 30-45). Unclipped, SAITS can LEARN to occasionally
# emit those dropouts, producing aberrant fills. Clipping the standardized
# values to a sane band tames the artifacts while leaving normal data (|z| < 10
# always) untouched. NaNs are preserved by np.clip.
CLIP_Z = 10.0


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def _to_datetime(values: pd.Series) -> pd.Series:
    """Coerce a TIMESTAMP column that may mix Excel serial floats and native
    Python datetime objects (the case when openpyxl auto-parsed dates)."""
    out = pd.Series(pd.NaT, index=values.index, dtype="datetime64[ns]")
    is_dt = values.map(lambda v: isinstance(v, pd.Timestamp) or hasattr(v, "year"))
    if is_dt.any():
        out.loc[is_dt] = pd.to_datetime(values[is_dt].tolist())
    if (~is_dt).any():
        serials = pd.to_numeric(values[~is_dt], errors="coerce")
        out.loc[~is_dt] = EXCEL_EPOCH + pd.to_timedelta(serials, unit="D")
    return out


def _find_timestamp_header(rows, scan: int = 6):
    """Return (header_row_idx, ts_col_idx) of the first 'TIMESTAMP' cell within
    the first `scan` rows, or None if absent."""
    for i, row in enumerate(rows[:scan]):
        for j, cell in enumerate(row):
            if cell == "TIMESTAMP":
                return i, j
    return None


def _resolve_toa5_sheet(path: Path, sheet_name: str | None):
    """Pick which worksheet holds the TOA5 table. Processed workbooks often put
    a 'notes'/cover sheet first, so reading worksheets[0] blindly fails. Honor an
    explicit sheet_name; otherwise return the FIRST sheet whose first rows carry
    a TIMESTAMP header. Cheap: only peeks the first 6 rows of each sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet '{sheet_name}' not in {path.name} (sheets: {wb.sheetnames})")
            return sheet_name
        for ws in wb.worksheets:
            peek = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                peek.append(row)
                if i >= 5:
                    break
            if _find_timestamp_header(peek) is not None:
                return ws.title
        raise ValueError(f"TIMESTAMP header not found in any sheet of {path.name} (sheets: {wb.sheetnames})")
    finally:
        wb.close()


def load_toa5_xlsx(path: str | Path, cache: bool = True, sheet_name: str | None = None) -> pd.DataFrame:
    """Read a Campbell Scientific TOA5 xlsx into a DatetimeIndex'd DataFrame.

    Locates the real header by the literal "TIMESTAMP" cell (some files have
    stray empty leading rows/columns) and, crucially, searches ACROSS sheets:
    multi-sheet workbooks (e.g. a processing workbook whose sheet 0 is 'notes')
    used to fail here with "TIMESTAMP header not found". Pass `sheet_name` to
    force a sheet, else the first sheet carrying a TIMESTAMP header is used.

    Cache is pickle-gzip rather than parquet because pyarrow's OpenMP runtime
    conflicts with PyTorch on Windows when torch is imported first. The cache
    key includes the resolved sheet so two sheets of the same file don't clash.
    """
    path = Path(path)
    sheet = _resolve_toa5_sheet(path, sheet_name)
    cache_path = path.with_suffix(f".{sheet}.pkl.gz")

    if cache and cache_path.exists() and cache_path.stat().st_mtime >= path.stat().st_mtime:
        return pd.read_pickle(cache_path, compression="gzip")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    hdr = _find_timestamp_header(rows)
    if hdr is None:
        raise ValueError(f"TIMESTAMP header not found in first 6 rows of sheet '{sheet}' in {path.name}")
    header_idx, ts_col_idx = hdr

    headers = [c for c in rows[header_idx][ts_col_idx:] if c is not None]
    n_cols = len(headers)

    data_rows = []
    for row in rows[header_idx + 3:]:  # skip header + units + aggregation rows
        slice_ = row[ts_col_idx : ts_col_idx + n_cols]
        if all(v is None for v in slice_):
            continue
        data_rows.append(slice_)

    df = pd.DataFrame(data_rows, columns=headers)
    df["TIMESTAMP"] = _to_datetime(df["TIMESTAMP"])
    for c in df.columns[1:]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df = df.dropna(subset=["TIMESTAMP"]).set_index("TIMESTAMP").sort_index()

    if cache:
        df.to_pickle(cache_path, compression="gzip")
    return df


def load_generic_table(path: str | Path) -> pd.DataFrame:
    """Load a generic in-memory dataset exported by the Rust side (Parquet or
    CSV) into a DatetimeIndex'd DataFrame. Used to train on CALCULATED data
    (e.g. T600) rather than raw TOA5 xlsx — the columns are arbitrary, not the
    TC_/SF_ naming convention, so the caller passes feature_columns explicitly.
    """
    path = Path(path)
    if path.suffix.lower() in (".parquet", ".pq"):
        df = pd.read_parquet(path)
    else:
        # Rust's CSV export uses ';' — fall back to ',' if that yields 1 col.
        df = pd.read_csv(path, sep=";")
        if df.shape[1] == 1:
            df = pd.read_csv(path, sep=",")

    ts_col = next((c for c in df.columns if str(c).upper() == "TIMESTAMP"), None)
    if ts_col is None:
        raise ValueError("Aucune colonne TIMESTAMP dans le dataset à entraîner.")
    df[ts_col] = _to_datetime(df[ts_col])
    for c in df.columns:
        if c == ts_col:
            continue
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df = df.dropna(subset=[ts_col]).set_index(ts_col).sort_index()
    # Collapse accidental duplicate timestamps (keep first non-null).
    if df.index.has_duplicates:
        df = df.groupby(df.index).first().sort_index()
    return df


def load_any(path: str | Path) -> pd.DataFrame:
    """Format-aware loader: Parquet/CSV (exported by the Rust side, e.g. a
    calculated dataset or env data) go through the generic loader; raw xlsx go
    through the TOA5-aware loader. openpyxl only reads .xlsx, so this avoids the
    "openpyxl does not support .parquet" error during training/detection."""
    p = Path(path)
    if p.suffix.lower() in (".parquet", ".pq", ".csv"):
        return load_generic_table(p)
    return load_toa5_xlsx(p)


def load_env_file(path: str | Path) -> pd.DataFrame:
    """Load environmental data (VPD, PAR, ETo, …) into a DatetimeIndex'd
    DataFrame (any format)."""
    return load_any(path)


def align_env_to_grid(
    env_df: pd.DataFrame,
    target_index: pd.DatetimeIndex,
    columns: list[str] | None = None,
    step: str = "5min",
) -> pd.DataFrame:
    """Bring env data onto the sensor's regular grid.

    Strategy:
      - Resample env to ``step`` taking the mean per bucket (works for sub-step
        env: keeps the mean; for super-step env: forward-fills to fill the gap).
      - Reindex to ``target_index`` with forward-fill (limit = 1 day = 288 buckets
        at 5min) so a daily env value broadcasts to the whole day, but a stale
        env entry doesn't leak forever.

    Args:
        env_df: DataFrame with DatetimeIndex.
        target_index: the sensor grid timestamps (e.g. df.index after resample).
        columns: subset of env columns to keep. None = all.
        step: target sampling step, must match the sensor grid.
    """
    if columns is not None:
        missing = [c for c in columns if c not in env_df.columns]
        if missing:
            raise ValueError(f"Env file missing columns: {missing}")
        env_df = env_df[columns]
    resampled = env_df.resample(pd.Timedelta(step), label="left", closed="left").mean()
    # Reindex onto the target grid; forward-fill up to 1 day so daily ETP
    # broadcasts to its day but a sensor gap doesn't carry an arbitrary stale
    # env value forever.
    limit = max(1, int(pd.Timedelta("1D") / pd.Timedelta(step)))
    aligned = resampled.reindex(target_index).ffill(limit=limit)
    return aligned


# Names accepted in the `time_features` param; cheap to compute, easy to name.
_TIME_FEATURE_NAMES = ("doy", "doy_sin", "doy_cos", "hod_sin", "hod_cos", "tabs")


def build_time_features(index: pd.DatetimeIndex, names: list[str]) -> pd.DataFrame:
    """Engineered temporal features. We give the model cyclic encodings rather
    than raw day/hour numbers — SAITS handles linear features fine but sin/cos
    pairs let it learn the periodic structure with fewer parameters.

    Available names:
      doy        — raw day-of-year (1..366), normalised to [0, 1]
      doy_sin    — sin(2π·doy/365.25)
      doy_cos    — cos(2π·doy/365.25)
      hod_sin    — sin(2π·hour_frac/24)
      hod_cos    — cos(2π·hour_frac/24)
      tabs       — ABSOLUTE calendar position, (year-2000 + doy/365.25)/50.

    `tabs` exists because every other feature here is periodic: doy is the same
    number on 15 March 2020 and 15 March 2023, so a model fed only these cannot
    tell one year from another. Sensor swaps and recalibrations shift a series'
    level between years, and normalisation is global (one mu/sigma per column
    for the whole training span) — so without an "which era is this" input the
    model reconstructs every gap towards the all-years mean, sitting visibly
    below the local level of a high year.

    Deliberately absolute rather than normalised over the data span: inference
    rebuilds these features from the prediction index alone, with no memory of
    the training range, so a span-relative encoding would silently mean
    something different at predict time.
    """
    unknown = [n for n in names if n not in _TIME_FEATURE_NAMES]
    if unknown:
        raise ValueError(f"Unknown time features: {unknown}. Valid: {_TIME_FEATURE_NAMES}")
    out = pd.DataFrame(index=index)
    doy = index.dayofyear.values.astype("float64")
    hod = index.hour.values.astype("float64") + index.minute.values.astype("float64") / 60.0
    twopi = 2.0 * np.pi
    for n in names:
        if n == "doy":     out[n] = doy / 365.25
        if n == "doy_sin": out[n] = np.sin(twopi * doy / 365.25)
        if n == "doy_cos": out[n] = np.cos(twopi * doy / 365.25)
        if n == "hod_sin": out[n] = np.sin(twopi * hod / 24.0)
        if n == "hod_cos": out[n] = np.cos(twopi * hod / 24.0)
        if n == "tabs":
            year = index.year.values.astype("float64")
            out[n] = (year - 2000.0 + doy / 365.25) / 50.0
    return out


def load_and_concat(paths: list[str | Path]) -> pd.DataFrame:
    """Load all files and combine them on the UNION of their columns.

    Files from different trees / stations rarely share TC_/SF_ columns
    (e.g. Niakhar 1 vs Niakhar 2). Restricting to the intersection would
    drop every sensor — instead we keep the full set and let cells stay
    NaN where a file doesn't cover that sensor. SAITS handles missing
    cells natively through its mask, so this is the right merge.

    When two files share a timestamp (overlap), we collapse with
    ``groupby.first()`` which picks the first non-null value per column.
    """
    if not paths:
        raise ValueError("No files provided.")
    dfs = [load_any(p) for p in paths]
    if len(dfs) == 1:
        return dfs[0]
    # axis=0 with default join='outer' gives the union of columns; rows from
    # a given file have NaN for columns it doesn't carry. Then collapse on
    # timestamp to merge overlapping periods.
    out = pd.concat(dfs, axis=0)
    out = out.groupby(out.index).first().sort_index()
    return out


# ---------------------------------------------------------------------------
# Sensor naming convention
# ---------------------------------------------------------------------------
# TC_<tree>-<branch>-<num>-<side>     e.g. TC_1a-RP-1-am
# SF_<tree>-<branch>-<num>            e.g. SF_1a-TA-1

_TC_RE = re.compile(r"^TC_([0-9]+[a-z])-([A-Z]+)-(\d+)-(am|av)$")
_SF_RE = re.compile(r"^SF_([0-9]+[a-z])-([A-Z]+)-(\d+)$")


@dataclass(frozen=True)
class Sensor:
    name: str
    type: str        # 'TC' or 'SF'
    tree: str
    branch: str
    num: int
    side: str | None  # 'am'/'av' for TC, None for SF


def parse_sensors(columns: list[str]) -> list[Sensor]:
    sensors: list[Sensor] = []
    for c in columns:
        m = _TC_RE.match(c)
        if m:
            sensors.append(Sensor(c, "TC", m.group(1), m.group(2), int(m.group(3)), m.group(4)))
            continue
        m = _SF_RE.match(c)
        if m:
            sensors.append(Sensor(c, "SF", m.group(1), m.group(2), int(m.group(3)), None))
    return sensors


# ---------------------------------------------------------------------------
# Resampling, segments, windowing
# ---------------------------------------------------------------------------

def resample_to_grid(df: pd.DataFrame, step: str = "5min") -> pd.DataFrame:
    return df.resample(pd.Timedelta(step), label="left", closed="left").mean()


def find_continuous_segments(
    df: pd.DataFrame,
    max_gap: str = "6h",
) -> list[tuple[pd.Timestamp, pd.Timestamp]]:
    """Return [(start, end)] for periods with no gap > max_gap. A gap is
    consecutive rows that are entirely NaN."""
    max_gap_td = pd.Timedelta(max_gap)
    has_data = df.notna().any(axis=1)
    if not has_data.any():
        return []
    valid_idx = df.index[has_data]
    deltas = valid_idx.to_series().diff()
    breaks = (deltas > max_gap_td).to_numpy().copy()
    breaks[0] = True

    segments = []
    seg_start = None
    prev_ts = None
    for i, ts in enumerate(valid_idx):
        if breaks[i]:
            if seg_start is not None:
                segments.append((seg_start, prev_ts))
            seg_start = ts
        prev_ts = ts
    if seg_start is not None:
        segments.append((seg_start, prev_ts))
    return segments


def make_windows_segmented(
    df_sub: pd.DataFrame,
    feat_cols: list[str],
    mu: pd.Series,
    sig: pd.Series,
    window: int,
    stride: int,
    max_nan_frac: float = 0.5,
    max_gap: str = "6h",
) -> np.ndarray:
    """Build (n_windows, window, n_features) z-scored arrays, never crossing
    a continuity gap. Skips windows with too much missingness."""
    segs = find_continuous_segments(df_sub, max_gap=max_gap)
    out = []
    for s, e in segs:
        seg = df_sub.loc[s:e]
        if len(seg) < window:
            continue
        X = np.clip(((seg[feat_cols] - mu) / sig).values, -CLIP_Z, CLIP_Z)
        for i in range(0, X.shape[0] - window + 1, stride):
            w = X[i:i + window]
            if np.isnan(w).mean() < max_nan_frac:
                out.append(w)
    return np.stack(out) if out else np.empty((0, window, len(feat_cols)))
